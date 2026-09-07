from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


HERE = Path(__file__).resolve().parent
DOCX = HERE / "ReSight-Afeka-POC-Proposal.docx"


def set_rfonts(run, name):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:cs"), name)


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)


def find_style(styles, name):
    try:
        return styles[name]
    except KeyError:
        for style in styles:
            if style.name == name:
                return style
        raise


def add_bottom_border(paragraph, color="D9E2EC", size="8"):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)


def add_shading(paragraph, fill):
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)


doc = Document(DOCX)

# Remove Pandoc's import of the browser <title>; the proposal already has its own title.
if doc.paragraphs and doc.paragraphs[0].text.strip() == "ReSight Afeka POC Prescreening Proposal":
    remove_paragraph(doc.paragraphs[0])

section = doc.sections[0]
section.start_type = WD_SECTION.NEW_PAGE
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(10.2)
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal._element.rPr.rFonts.set(qn("w:cs"), "Arial")
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.15

for style_name, size, color, before, after in [
    ("Heading 1", 18, "12355B", 0, 4),
    ("Heading 2", 12.5, "0F4C81", 10, 4),
]:
    style = find_style(styles, style_name)
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = True
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

caption_style = styles["Caption"]
caption_style.font.name = "Arial"
caption_style.font.size = Pt(8.5)
caption_style.font.color.rgb = RGBColor(82, 96, 109)
caption_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
caption_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
caption_style._element.rPr.rFonts.set(qn("w:cs"), "Arial")

for idx, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if not text:
        continue
    if idx == 0 and text.startswith("Preliminary proposal"):
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in para.runs:
            set_rfonts(run, "Arial")
            run.font.size = Pt(9.2)
            run.font.color.rgb = RGBColor(82, 96, 109)
        continue
    if text.startswith("ReSight: AI for Repurposing"):
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in para.runs:
            set_rfonts(run, "Arial")
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(18, 53, 91)
        add_bottom_border(para)
        continue
    if text.startswith("ReSight: מערכת"):
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ppr = para._p.get_or_add_pPr()
        bidi = ppr.find(qn("w:bidi"))
        if bidi is None:
            bidi = OxmlElement("w:bidi")
            ppr.append(bidi)
        for run in para.runs:
            set_rfonts(run, "Arial")
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(36, 59, 83)
        continue
    if text in {"תקציר בעברית"}:
        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_shading(para, "F8FBFF")
    if text.startswith("Applicant:") or text.startswith("Email:") or text.startswith("Domain:"):
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        para.paragraph_format.space_after = Pt(1)
    elif para.style.name not in {"Heading 1", "Heading 2", "Caption"}:
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        para.paragraph_format.space_after = Pt(5)
        para.paragraph_format.line_spacing = 1.15
    for run in para.runs:
        set_rfonts(run, "Arial")

# Keep figures in page width.
max_width = Inches(6.9)
for shape in doc.inline_shapes:
    if shape.width > max_width:
        ratio = max_width / shape.width
        shape.width = int(max_width)
        shape.height = int(shape.height * ratio)

doc.save(DOCX)
print(DOCX)
